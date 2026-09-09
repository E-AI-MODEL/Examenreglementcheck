import tempfile,unittest,sys,json
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from app.parser import parse_document
from app.registers import build_registers
from app.analyzer import analyze_document
from app.server import app

class V31RegressionTests(unittest.TestCase):
 def doc(self,td,paras,year='2026-2027'):
  p=Path(td)/'x.docx';d=Document()
  for x in paras:d.add_paragraph(x)
  d.save(p);return parse_document(p,school_year=year,school_types=['vwo'])
 def test_year_and_date_are_not_articles(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['2026-2027','1 januari 2027 wordt het reglement van kracht.','Artikel 2.1 Begrippen','2.2 Procedure'])
   self.assertEqual([u['article'] for u in doc['units'] if u['type']=='article'],['2.1','2.2'])
 def test_advisor_and_decision_maker_are_not_role_conflict(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['De examencommissie adviseert over maatregelen en de rector besluit over maatregelen.'])
   result=analyze_document(doc,root=ROOT,run_id='x')
   self.assertFalse(any('role_conflict' in f['finding_type'] for f in result['findings']))
   regs=result['registers']['procedure_registry'];self.assertTrue(any(p['action']=='advise' and p['actor']=='examencommissie' for p in regs));self.assertTrue(any(p['action']=='decide' and p['actor']=='rector' for p in regs))
 def test_filing_term_and_decision_term_are_not_conflict(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['De kandidaat kan binnen 5 dagen beroep instellen.','De commissie van beroep beslist binnen 2 weken op het beroep.'])
   result=analyze_document(doc,root=ROOT,run_id='x')
   self.assertFalse(any('term_conflict' in f['finding_type'] for f in result['findings']))
   kinds={x['deadline_type'] for x in result['registers']['term_registry']};self.assertIn('filing',kinds);self.assertIn('decision',kinds)
 def test_same_filing_step_different_values_is_conflict(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['De kandidaat kan binnen 2 dagen beroep instellen.','De kandidaat kan binnen 5 dagen beroep instellen.'])
   result=analyze_document(doc,root=ROOT,run_id='x')
   self.assertTrue(any('term_conflict' in f['finding_type'] for f in result['findings']))
 def test_generic_word_wet_does_not_hide_internal_broken_reference(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['Artikel 1.1 Algemeen','Volgens de wet geldt voor de interne procedure artikel 12.3.'])
   regs=build_registers(doc);ref=next(r for r in regs['cross_reference_registry'] if r['target']=='12.3')
   self.assertEqual(ref['reference_scope'],'internal');self.assertFalse(ref['resolved'])
 def test_named_wvo_reference_is_external(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['Artikel 1.1 Volgens artikel 2.60 WVO 2020 wordt het examenreglement vastgesteld.'])
   ref=build_registers(doc)['cross_reference_registry'][0];self.assertEqual(ref['reference_scope'],'external');self.assertEqual(ref['target_source'],'wvo_2020');self.assertTrue(ref['resolved'])
 def test_confirmed_year_recomputes_exam_year_atomically(self):
  c=TestClient(app)
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'x.docx';d=Document();d.add_paragraph('Examenreglement');d.add_paragraph('De kandidaat kan binnen 5 dagen beroep instellen.');d.save(p)
   r=c.post('/api/parse',files={'file':('x.docx',p.read_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'school_year':'2025-2026','school_types':'vwo','document_status':'concept','ai_mode':'off'});self.assertEqual(r.status_code,200);rid=r.json()['run_id']
   r=c.post('/api/analyze/'+rid,json={'school_year':'2026-2027','school_types':['vwo'],'document_status':'concept','confirmed':True,'ai_mode':'off'});self.assertEqual(r.status_code,200);x=r.json()
   self.assertEqual(x['school_year'],'2026-2027');self.assertEqual(x['exam_year'],2027);self.assertEqual(x['document']['metadata']['expected_exam_year'],2027);self.assertEqual(x['confirmed_context']['exam_year'],2027)
 def test_source_gate_is_executed_but_semantic_validator_still_blocks_must(self):
  with tempfile.TemporaryDirectory() as td:
   doc=self.doc(td,['Artikel 1.1 De kandidaat kan binnen 5 dagen beroep instellen.'])
   result=analyze_document(doc,root=ROOT,run_id='x')
   self.assertGreater(result['coverage']['source_gate_checked'],0);self.assertEqual(result['coverage']['source_gate_checked'],len(result['source_candidates']));self.assertEqual(result['coverage']['semantic_evidence_validator'],'not_implemented');self.assertFalse(result['coverage']['legal_must_enabled'])
   for c in result['source_candidates']:self.assertIn('source_gate',c)
if __name__=='__main__':unittest.main(verbosity=2)
