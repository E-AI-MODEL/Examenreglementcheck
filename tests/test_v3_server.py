import tempfile,unittest,sys,json
from pathlib import Path
from docx import Document
from fastapi.testclient import TestClient
from jsonschema import Draft202012Validator
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from app.server import app
class ServerTests(unittest.TestCase):
 def fixture(self,td):
  p=Path(td)/'x.docx';d=Document();d.add_paragraph('Examenreglement 2026-2027');d.add_paragraph('De kandidaat meldt zich binnen 2 dagen aan voor beroep.');d.add_paragraph('De kandidaat meldt zich binnen 5 dagen aan voor beroep.');d.save(p);return p
 def test_health_and_ai_block(self):
  c=TestClient(app);h=c.get('/api/health').json();self.assertTrue(h['ok']);self.assertEqual(h['version'],'0.4.1');self.assertFalse(h['ai']['available'])
  with tempfile.TemporaryDirectory() as td:
   p=self.fixture(td);r=c.post('/api/parse',files={'file':('x.docx',p.read_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'school_year':'2026-2027','ai_mode':'on'});self.assertEqual(r.status_code,409)
 def test_parse_confirm_analyze_run_contract(self):
  c=TestClient(app)
  with tempfile.TemporaryDirectory() as td:
   p=self.fixture(td);r=c.post('/api/parse',files={'file':('x.docx',p.read_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'school_year':'2026-2027','school_types':'vwo','document_status':'concept','ai_mode':'off'});self.assertEqual(r.status_code,200);x=r.json();self.assertEqual(x['status'],'awaiting_confirmation')
   rid=x['run_id'];r=c.post('/api/analyze/'+rid,json={'school_year':'2026-2027','school_types':['vwo'],'document_status':'concept','confirmed':True,'ai_mode':'off'});self.assertEqual(r.status_code,200);x=r.json();self.assertEqual(x['status'],'partial');self.assertTrue(x['findings']);self.assertEqual(x['phase_status']['B_volledigheid'],'not_implemented');self.assertEqual(x['phase_status']['I_evidence_validatie'],'partial');Draft202012Validator(json.loads((ROOT/'schemas/run.schema.json').read_text())).validate(x)
 def test_fake_docx_is_rejected_by_content(self):
  c=TestClient(app);r=c.post('/api/parse',files={'file':('x.docx',b'not a zip','application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'school_year':'2026-2027','school_types':'vwo','ai_mode':'off'});self.assertEqual(r.status_code,415)
 def test_fake_pdf_is_rejected_by_content(self):
  c=TestClient(app);r=c.post('/api/parse',files={'file':('x.pdf',b'not pdf','application/pdf')},data={'school_year':'2026-2027','school_types':'vwo','ai_mode':'off'});self.assertEqual(r.status_code,415)
class ReviewApiTests(unittest.TestCase):
 def test_review_persists_change_set_and_delete_run(self):
  c=TestClient(app)
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'r.docx';d=Document();d.add_paragraph('Artikel 1.1');d.add_paragraph('De kandidaat kan binnen 2 dagen beroep instellen.');d.add_paragraph('De kandidaat kan binnen 5 dagen beroep instellen.');d.save(p)
   x=c.post('/api/parse',files={'file':('r.docx',p.read_bytes(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document')},data={'school_year':'2026-2027','school_types':'vwo','document_status':'concept','ai_mode':'off'}).json();rid=x['run_id']
   x=c.post('/api/analyze/'+rid,json={'school_year':'2026-2027','school_types':['vwo'],'document_status':'concept','confirmed':True,'ai_mode':'off'}).json();fid=x['findings'][0]['finding_id']
   r=c.patch(f'/api/runs/{rid}/findings/{fid}',json={'human_status':'accepted','review_notes':'Controle door commissie','in_change_set':True,'suggested_text':'Nieuwe tekst'});self.assertEqual(r.status_code,200);f=r.json();self.assertEqual(f['human_status'],'accepted');self.assertTrue(f['in_change_set'])
   ch=c.get(f'/api/runs/{rid}/changes').json()['changes'];self.assertEqual(len(ch),1);self.assertEqual(ch[0]['finding_id'],fid)
   self.assertTrue(c.delete(f'/api/runs/{rid}').json()['deleted']);self.assertEqual(c.get(f'/api/runs/{rid}').status_code,404)
if __name__=='__main__':unittest.main()
