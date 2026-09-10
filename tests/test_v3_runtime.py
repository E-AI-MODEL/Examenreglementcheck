import json,tempfile,unittest
from pathlib import Path
from docx import Document
import pymupdf as fitz
from jsonschema import Draft202012Validator,RefResolver
import sys
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from app.parser import parse_document
from app.registers import build_registers
from app.analyzer import analyze_document
from app.ai_adapter import status,analyze as ai_analyze

class V3RuntimeTests(unittest.TestCase):
 def make_docx(self,p):
  d=Document();d.add_heading('Examenreglement 2026-2027',0);d.add_heading('Artikel 1 Algemeen',1);d.add_paragraph('Artikel 1.1 Begrippen');d.add_paragraph('De kandidaat meldt zich binnen 2 dagen aan voor beroep.');d.add_paragraph('De kandidaat meldt zich binnen 5 dagen aan voor beroep.');d.add_paragraph('Artikel 2 Maatregelen');d.add_paragraph('De rector besluit over maatregelen bij onregelmatigheden.');d.add_paragraph('De examencommissie besluit over maatregelen bij onregelmatigheden.');d.add_paragraph('Voor de procedure wordt verwezen naar artikel 19.4.');d.add_paragraph('commissie van beroep betekent de commissie die het beroep behandelt.');d.add_paragraph('commissie van beroep betekent het orgaan dat maatregelen vaststelt.');d.add_table(rows=2,cols=2).cell(0,0).text='Tabel';d.save(p)
 def test_docx_parse_schema_and_registers(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'r.docx';self.make_docx(p);doc=parse_document(p,school_year='2026-2027',school_types=['vwo'])
   schema=json.loads((ROOT/'schemas/document.schema.json').read_text());resolver=RefResolver(base_uri=(ROOT/'schemas/').as_uri()+'/',referrer=schema);Draft202012Validator(schema,resolver=resolver).validate(doc)
   regs=build_registers(doc);Draft202012Validator(json.loads((ROOT/'schemas/registers.schema.json').read_text())).validate(regs)
   self.assertGreaterEqual(len(regs['term_registry']),2);self.assertTrue(any(not x['resolved'] for x in regs['cross_reference_registry']))
 def test_pdf_parse_keeps_page_anchors(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'r.pdf';pdf=fitz.open();page=pdf.new_page();page.insert_text((72,72),'Examenreglement 2026-2027\nArtikel 1.1 De kandidaat kan binnen 5 dagen beroep instellen.');pdf.save(p)
   doc=parse_document(p,school_year='2026-2027');self.assertTrue(doc['units']);self.assertEqual(doc['units'][0]['page'],1);self.assertTrue(doc['units'][0]['anchor_id'].startswith('pdf-p001'))
 def test_analyzer_finds_internal_conflicts_without_must(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'r.docx';self.make_docx(p);doc=parse_document(p,school_year='2026-2027')
   result=analyze_document(doc,root=ROOT,run_id='run-test');sevs={f['severity'] for f in result['findings']}
   self.assertIn('internal_conflict',sevs);self.assertIn('editorial',sevs);self.assertNotIn('must',sevs);self.assertFalse(result['coverage']['legal_must_enabled'])
 def test_external_legal_reference_not_marked_broken(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'r.docx';d=Document();d.add_paragraph('Artikel 1.1 Volgens artikel 2.60 WVO 2020 wordt het examenreglement vastgesteld.');d.save(p)
   regs=build_registers(parse_document(p));self.assertTrue(regs['cross_reference_registry'][0]['resolved']);self.assertTrue(regs['cross_reference_registry'][0]['external'])
 def test_ai_is_explicitly_unavailable(self):
  self.assertFalse(status().available)
  with self.assertRaises(RuntimeError):ai_analyze('test')
if __name__=='__main__':unittest.main(verbosity=2)
