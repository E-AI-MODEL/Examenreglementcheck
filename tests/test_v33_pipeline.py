import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document
from fastapi.testclient import TestClient

import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app.analyzer import analyze_document
from app.counter_review import run_counter_review
from app.parser import parse_document
from app.scope import rule_in_scope
from app.server import app


class V33PipelineTests(unittest.TestCase):
    def make_docx(self, path: Path) -> Path:
        document = Document()
        document.add_heading("Examenreglement 2026-2027", 0)
        document.add_paragraph("De rector besluit over maatregelen bij onregelmatigheden.")
        document.add_paragraph("De kandidaat kan binnen 2 dagen beroep instellen.")
        document.add_paragraph("De kandidaat kan binnen 5 dagen beroep instellen.")
        document.save(path)
        return path

    def test_full_in_scope_pipeline_and_explicit_exclusions(self):
        with tempfile.TemporaryDirectory() as td:
            parsed = parse_document(
                self.make_docx(Path(td) / "reglement.docx"),
                school_year="2026-2027",
                school_types=["vwo"],
            )
            result = analyze_document(parsed, root=ROOT, run_id="v33", as_of="2026-09-10")
        for key in ("completeness", "actuality", "comparisons", "counter_review", "evidence_validation"):
            self.assertEqual(result[key]["status"], "complete")
        self.assertFalse(result["coverage"]["legal_must_enabled"])
        self.assertFalse(any(finding["severity"] == "must" for finding in result["findings"]))
        excluded_prefixes = ("SE-", "CE-", "PTA-", "HUL-", "ROO-")
        self.assertFalse(any(row["rule_id"].startswith(excluded_prefixes) for row in result["source_candidates"]))
        checked_sources = {row["source_id"] for row in result["actuality"]["sources"]}
        completeness_sources = {row["source_id"] for row in result["completeness"]["checks"]}
        self.assertTrue(completeness_sources <= checked_sources)
        self.assertTrue(all(row["legal_evidence"] is False for row in result["comparisons"]["comparisons"]))

    def test_scope_also_excludes_se_content_hidden_behind_generic_prefix(self):
        self.assertFalse(rule_in_scope({"rule_id": "REG-999", "topic": "Inhalen", "rule_summary": "Regelt het schoolexamen en SE."}))
        self.assertTrue(rule_in_scope({"rule_id": "REG-998", "topic": "Bekendmaking", "rule_summary": "Het reglement wordt gepubliceerd."}))

    def test_ocr_fallback_creates_anchored_units(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "scan.pdf"
            pdf = pymupdf.open()
            pdf.new_page()
            pdf.save(path)
            with patch("app.parser._ocr_page_text", return_value="Examenreglement 2026-2027\nArtikel 1.1 Algemene bepalingen voor de kandidaat"):
                parsed = parse_document(path, school_year="2026-2027", school_types=["vwo"])
        self.assertEqual(parsed["parsing_status"], "complete")
        self.assertTrue(parsed["units"])
        self.assertEqual(parsed["units"][0]["extraction_method"], "ocr_tesseract")
        self.assertTrue(parsed["units"][0]["anchor_id"].startswith("pdf-p001-ocr"))
        self.assertEqual(parsed["warnings"][0]["code"], "page_ocr_used")

    def test_counter_review_downgrades_unsupported_must(self):
        findings = [{"finding_id": "f-1", "claim": "Verplichte regel", "severity": "must", "confidence": "high", "finding_type": ["legal"], "evidence": []}]
        result = run_counter_review(findings, parsing_status="complete")
        self.assertEqual(findings[0]["severity"], "human_review")
        self.assertIn("unsupported_must_downgraded", findings[0]["counter_review"]["checks"])
        self.assertEqual(result["changed_findings"], 1)

    def test_run_resume_export_delete_and_empty_school_type_validation(self):
        client = TestClient(app)
        with tempfile.TemporaryDirectory() as td:
            path = self.make_docx(Path(td) / "api.docx")
            parsed = client.post(
                "/api/parse",
                files={"file": ("api.docx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
                data={"school_year": "2026-2027", "school_types": "vwo", "document_status": "concept", "ai_mode": "off"},
            ).json()
        run_id = parsed["run_id"]
        try:
            rejected = client.post(
                f"/api/analyze/{run_id}",
                json={"school_year": "2026-2027", "school_types": [], "document_status": "concept", "confirmed": True, "ai_mode": "off"},
            )
            self.assertEqual(rejected.status_code, 422)
            analyzed = client.post(
                f"/api/analyze/{run_id}",
                json={"school_year": "2026-2027", "school_types": ["vwo"], "document_status": "concept", "confirmed": True, "ai_mode": "off"},
            )
            self.assertEqual(analyzed.status_code, 200)
            self.assertEqual(analyzed.json()["status"], "complete")
            self.assertIn(run_id, {row["run_id"] for row in client.get("/api/runs").json()["runs"]})
            self.assertEqual(client.get(f"/api/runs/{run_id}").json()["run_id"], run_id)
            export = client.get(f"/api/runs/{run_id}/export")
            self.assertEqual(export.status_code, 200)
            self.assertEqual(export.json()["run_id"], run_id)
            markdown = client.get(f"/api/runs/{run_id}/export.md")
            self.assertEqual(markdown.status_code, 200)
            self.assertIn("# Examenreglement-check — rapport", markdown.text)
            self.assertIn("attachment", markdown.headers["content-disposition"])
        finally:
            client.delete(f"/api/runs/{run_id}")


if __name__ == "__main__":
    unittest.main(verbosity=2)
